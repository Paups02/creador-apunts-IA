"""
DOCX formatter module - Deep Learning Notes format
Creates beautifully formatted hierarchical notes with comprehensive explanations
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from typing import Dict, List
from pathlib import Path

from .config import (
    Colors, FontSizes, FontNames, LINE_SPACING,
    SPACE_AFTER, SPACE_BEFORE, IMAGE_WIDTH_INCHES
)


class DeepLearningFormatter:
    """Create comprehensive Deep Learning Notes formatted DOCX documents"""

    def __init__(self):
        """Initialize formatter"""
        self.doc = Document()
        self._setup_document_margins()

    def _setup_document_margins(self):
        """Set up document margins for better layout"""
        sections = self.doc.sections
        for section in sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

    def create_learning_notes(self, notes_data: Dict, output_path: str):
        """
        Create Deep Learning Notes document

        Args:
            notes_data: Structured notes data from AI
            output_path: Where to save the DOCX file
        """
        # Add title
        self._add_title(notes_data.get("title", "Learning Notes"))

        # Add sections
        sections = notes_data.get("sections", [])
        for section_idx, section in enumerate(sections):
            self._add_section(section, section_idx)

            # Add spacing between sections
            if section_idx < len(sections) - 1:
                self.doc.add_paragraph()

        # Add key takeaways at the end
        if notes_data.get("key_takeaways"):
            self._add_key_takeaways(notes_data["key_takeaways"])

        # Save document
        output_file = Path(output_path)
        output_file.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(output_file))

    def _add_title(self, title: str):
        """Add formatted title to document"""
        title_para = self.doc.add_paragraph()
        title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        title_run = title_para.add_run(title)
        title_run.font.size = FontSizes.TITLE
        title_run.font.name = FontNames.PRIMARY
        title_run.font.color.rgb = Colors.TITLE
        title_run.bold = True

        title_para.space_after = Pt(20)

        # Add decorative line
        line_para = self.doc.add_paragraph("_" * 80)
        line_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        line_para.space_after = Pt(15)

    def _add_section(self, section: Dict, section_idx: int):
        """Add a section with hierarchical content"""
        # Section heading
        heading = section.get("heading", f"Section {section_idx + 1}")
        heading_para = self.doc.add_paragraph()
        heading_run = heading_para.add_run(heading)
        heading_run.font.size = Pt(16)
        heading_run.font.name = FontNames.PRIMARY
        heading_run.font.color.rgb = Colors.HEADING
        heading_run.bold = True
        heading_para.space_after = Pt(10)

        # Content items
        content_items = section.get("content", [])
        for item in content_items:
            self._add_content_item(item)

    def _add_content_item(self, item: Dict):
        """Add a single content item with appropriate formatting"""
        item_type = item.get("type", "text")
        text = item.get("text", "")
        level = item.get("level", 0)

        if item_type == "intro":
            self._add_intro(text, level)
        elif item_type == "heading":
            self._add_heading_item(text, level)
        elif item_type == "explanation":
            self._add_explanation(text, level)
        elif item_type == "definition":
            self._add_definition(text, level)
        elif item_type == "example":
            self._add_example(text, level)
        elif item_type == "analogy":
            self._add_analogy(text, level)
        elif item_type == "bullet":
            self._add_bullet_item(text, level)
        elif item_type == "connection":
            self._add_connection(text, level)
        elif item_type == "warning":
            self._add_warning(text, level)
        elif item_type == "tldr":
            self._add_tldr(text, level)
        elif item_type == "table":
            self._add_table_item(item)
        elif item_type == "image":
            self._add_image_item(item, level)
        else:  # text
            self._add_text_item(text, level)

    def _add_intro(self, text: str, level: int):
        """Add introduction text"""
        para = self.doc.add_paragraph()
        run = para.add_run("📖 " + text)
        run.font.size = Pt(11)
        run.font.name = FontNames.PRIMARY
        run.font.italic = True
        run.font.color.rgb = RGBColor(50, 50, 50)
        para.paragraph_format.left_indent = Inches(0.25 * level)
        para.space_after = Pt(8)

    def _add_heading_item(self, text: str, level: int):
        """Add a subheading"""
        para = self.doc.add_paragraph()
        run = para.add_run(text)

        if level == 1:
            run.font.size = Pt(13)
            run.font.color.rgb = Colors.HEADING
        else:
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(100, 100, 100)

        run.font.name = FontNames.PRIMARY
        run.bold = True
        para.paragraph_format.left_indent = Inches(0.25 * level)
        para.space_after = Pt(6)

    def _add_explanation(self, text: str, level: int):
        """Add detailed explanation"""
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        run.font.size = FontSizes.MAIN_TEXT
        run.font.name = FontNames.PRIMARY
        para.paragraph_format.left_indent = Inches(0.25 * level)
        para.paragraph_format.line_spacing = LINE_SPACING
        para.space_after = Pt(8)

    def _add_definition(self, text: str, level: int):
        """Add definition with highlighting"""
        para = self.doc.add_paragraph()

        # Split on colon to separate term and definition
        if ":" in text:
            term, definition = text.split(":", 1)

            # Term in bold
            term_run = para.add_run("📝 " + term.strip() + ": ")
            term_run.font.size = FontSizes.MAIN_TEXT
            term_run.font.name = FontNames.PRIMARY
            term_run.bold = True
            term_run.font.color.rgb = RGBColor(0, 102, 204)  # Blue

            # Definition normal
            def_run = para.add_run(definition.strip())
            def_run.font.size = FontSizes.MAIN_TEXT
            def_run.font.name = FontNames.PRIMARY
        else:
            run = para.add_run("📝 " + text)
            run.font.size = FontSizes.MAIN_TEXT
            run.font.name = FontNames.PRIMARY
            run.bold = True
            run.font.color.rgb = RGBColor(0, 102, 204)

        # Light background shading
        pPr = para._element.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'E7F3FF')  # Very light blue
        pPr.append(shd)

        para.paragraph_format.left_indent = Inches(0.25 * level)
        para.paragraph_format.space_before = Pt(3)
        para.paragraph_format.space_after = Pt(3)

    def _add_example(self, text: str, level: int):
        """Add example with icon"""
        para = self.doc.add_paragraph()

        # Icon
        icon_run = para.add_run("💡 EXEMPLE: ")
        icon_run.font.size = FontSizes.MAIN_TEXT
        icon_run.font.name = FontNames.PRIMARY
        icon_run.bold = True
        icon_run.font.color.rgb = RGBColor(255, 153, 0)  # Orange

        # Example text
        text_run = para.add_run(text)
        text_run.font.size = FontSizes.MAIN_TEXT
        text_run.font.name = FontNames.PRIMARY
        text_run.font.italic = True

        # Light orange background
        pPr = para._element.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'FFF4E6')  # Very light orange
        pPr.append(shd)

        para.paragraph_format.left_indent = Inches(0.25 * level + 0.2)
        para.paragraph_format.space_after = Pt(8)

    def _add_analogy(self, text: str, level: int):
        """Add analogy/comparison"""
        para = self.doc.add_paragraph()

        icon_run = para.add_run("🔗 ANALOGIA: ")
        icon_run.font.size = FontSizes.MAIN_TEXT
        icon_run.font.name = FontNames.PRIMARY
        icon_run.bold = True
        icon_run.font.color.rgb = RGBColor(102, 51, 153)  # Purple

        text_run = para.add_run(text)
        text_run.font.size = FontSizes.MAIN_TEXT
        text_run.font.name = FontNames.PRIMARY

        para.paragraph_format.left_indent = Inches(0.25 * level + 0.2)
        para.space_after = Pt(8)

    def _add_bullet_item(self, text: str, level: int):
        """Add bullet point"""
        para = self.doc.add_paragraph(style='List Bullet')
        run = para.add_run(text)
        run.font.size = FontSizes.MAIN_TEXT
        run.font.name = FontNames.PRIMARY
        para.paragraph_format.left_indent = Inches(0.25 * level)
        para.space_after = Pt(4)

    def _add_connection(self, text: str, level: int):
        """Add connection to other concepts"""
        para = self.doc.add_paragraph()

        icon_run = para.add_run("🔄 CONNEXIÓ: ")
        icon_run.font.size = FontSizes.MAIN_TEXT
        icon_run.font.name = FontNames.PRIMARY
        icon_run.bold = True
        icon_run.font.color.rgb = RGBColor(0, 153, 76)  # Green

        text_run = para.add_run(text)
        text_run.font.size = FontSizes.MAIN_TEXT
        text_run.font.name = FontNames.PRIMARY

        para.paragraph_format.left_indent = Inches(0.25 * level)
        para.space_after = Pt(6)

    def _add_warning(self, text: str, level: int):
        """Add warning/common mistake"""
        para = self.doc.add_paragraph()

        icon_run = para.add_run("⚠️ ATENCIÓ: ")
        icon_run.font.size = FontSizes.MAIN_TEXT
        icon_run.font.name = FontNames.PRIMARY
        icon_run.bold = True
        icon_run.font.color.rgb = RGBColor(204, 0, 0)  # Red

        text_run = para.add_run(text)
        text_run.font.size = FontSizes.MAIN_TEXT
        text_run.font.name = FontNames.PRIMARY

        # Light red background
        pPr = para._element.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'FFE6E6')  # Very light red
        pPr.append(shd)

        para.paragraph_format.left_indent = Inches(0.25 * level + 0.2)
        para.space_after = Pt(8)

    def _add_tldr(self, text: str, level: int):
        """Add TL;DR summary"""
        para = self.doc.add_paragraph()

        icon_run = para.add_run("📌 RESUM RÀPID: ")
        icon_run.font.size = FontSizes.MAIN_TEXT
        icon_run.font.name = FontNames.PRIMARY
        icon_run.bold = True
        icon_run.font.color.rgb = Colors.HEADING

        text_run = para.add_run(text)
        text_run.font.size = FontSizes.MAIN_TEXT
        text_run.font.name = FontNames.PRIMARY
        text_run.bold = True

        # Light gray background
        pPr = para._element.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'F0F0F0')  # Light gray
        pPr.append(shd)

        para.paragraph_format.left_indent = Inches(0.25 * level)
        para.space_after = Pt(10)

    def _add_text_item(self, text: str, level: int):
        """Add regular text"""
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        run.font.size = FontSizes.MAIN_TEXT
        run.font.name = FontNames.PRIMARY
        para.paragraph_format.left_indent = Inches(0.25 * level)
        para.space_after = Pt(6)

    def _add_table_item(self, item: Dict):
        """Add a table"""
        headers = item.get("headers", [])
        rows = item.get("rows", [])

        if not headers or not rows:
            return

        # Add spacing before table
        self.doc.add_paragraph().space_after = Pt(6)

        # Create table
        num_cols = len(headers)
        num_rows = len(rows) + 1  # +1 for header

        table = self.doc.add_table(rows=num_rows, cols=num_cols)
        table.style = 'Light Grid Accent 1'

        # Add headers
        for col_idx, header in enumerate(headers):
            cell = table.rows[0].cells[col_idx]
            para = cell.paragraphs[0]
            run = para.add_run(str(header))
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = Colors.WHITE

            # Blue header background
            self._shade_cell(cell, Colors.HEADING)

        # Add data rows
        for row_idx, row_data in enumerate(rows, start=1):
            for col_idx, cell_data in enumerate(row_data):
                if col_idx < num_cols:
                    cell = table.rows[row_idx].cells[col_idx]
                    para = cell.paragraphs[0]
                    run = para.add_run(str(cell_data))
                    run.font.size = Pt(9)

        # Add spacing after table
        self.doc.add_paragraph().space_after = Pt(10)

    def _add_image_item(self, item: Dict, level: int):
        """Add a generated image with caption"""
        image_path = item.get("image_path")
        caption = item.get("caption", "")

        # Skip if image generation failed or no path
        if not image_path or item.get("image_failed"):
            # Add placeholder text if image failed
            if caption:
                para = self.doc.add_paragraph()
                run = para.add_run(f"🖼️ [Imatge no disponible: {caption}]")
                run.font.size = Pt(10)
                run.font.italic = True
                run.font.color.rgb = RGBColor(128, 128, 128)
                para.paragraph_format.left_indent = Inches(0.25 * level)
                para.space_after = Pt(8)
            return

        try:
            # Add spacing before image
            self.doc.add_paragraph().space_after = Pt(6)

            # Add the image
            para = self.doc.add_paragraph()
            para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run = para.add_run()
            run.add_picture(image_path, width=Inches(IMAGE_WIDTH_INCHES))

            # Add caption below image
            if caption:
                caption_para = self.doc.add_paragraph()
                caption_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

                caption_run = caption_para.add_run(f"📸 {caption}")
                caption_run.font.size = Pt(10)
                caption_run.font.name = FontNames.PRIMARY
                caption_run.font.italic = True
                caption_run.font.color.rgb = RGBColor(80, 80, 80)
                caption_para.space_after = Pt(12)

        except Exception as e:
            # If image insertion fails, add error message
            para = self.doc.add_paragraph()
            run = para.add_run(f"🖼️ [Error carregant imatge: {caption}]")
            run.font.size = Pt(10)
            run.font.italic = True
            run.font.color.rgb = RGBColor(128, 128, 128)
            para.paragraph_format.left_indent = Inches(0.25 * level)
            para.space_after = Pt(8)

    def _add_key_takeaways(self, takeaways: List[str]):
        """Add key takeaways section at the end"""
        # Add spacing and separator
        self.doc.add_paragraph()
        separator = self.doc.add_paragraph("═" * 80)
        separator.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        separator.space_after = Pt(10)

        # Title
        title_para = self.doc.add_paragraph()
        title_run = title_para.add_run("🎯 PUNTS CLAU")
        title_run.font.size = Pt(16)
        title_run.font.name = FontNames.PRIMARY
        title_run.font.color.rgb = Colors.TITLE
        title_run.bold = True
        title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        title_para.space_after = Pt(10)

        # Add takeaways
        for idx, takeaway in enumerate(takeaways, 1):
            para = self.doc.add_paragraph()

            # Number
            num_run = para.add_run(f"{idx}. ")
            num_run.font.size = Pt(12)
            num_run.font.name = FontNames.PRIMARY
            num_run.bold = True
            num_run.font.color.rgb = Colors.HEADING

            # Takeaway text
            text_run = para.add_run(takeaway)
            text_run.font.size = Pt(12)
            text_run.font.name = FontNames.PRIMARY

            # Light blue background
            pPr = para._element.get_or_add_pPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), 'E8F4F8')
            pPr.append(shd)

            para.paragraph_format.left_indent = Inches(0.5)
            para.space_after = Pt(8)

    def _shade_cell(self, cell, color: RGBColor):
        """Add background shading to a table cell"""
        cell_elem = cell._element
        cell_pr = cell_elem.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        hex_color = str(color)
        shd.set(qn('w:fill'), hex_color)
        cell_pr.append(shd)


def create_learning_notes_docx(notes_data: Dict, output_path: str, cleanup_temp_images: bool = True):
    """
    Convenience function to create Deep Learning Notes DOCX

    Args:
        notes_data: Structured notes data
        output_path: Output file path
        cleanup_temp_images: Whether to clean up temporary image files after creation
    """
    formatter = DeepLearningFormatter()
    formatter.create_learning_notes(notes_data, output_path)

    # Clean up temporary images if requested
    if cleanup_temp_images and "_image_generator" in notes_data:
        notes_data["_image_generator"].cleanup()
        del notes_data["_image_generator"]


# Alias for backwards compatibility
create_cornell_docx = create_learning_notes_docx
