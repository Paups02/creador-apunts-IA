"""
DOCX formatter module
Creates beautifully formatted DOCX files with Cornell Method layout
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from typing import Dict, List, Any
from pathlib import Path

from .config import (
    Colors, FontSizes, FontNames, LINE_SPACING,
    SPACE_AFTER, SPACE_BEFORE, CUE_COLUMN_WIDTH,
    NOTES_COLUMN_WIDTH, TABLE_TOTAL_WIDTH
)


class CornellFormatter:
    """Create Cornell Method formatted DOCX documents"""

    def __init__(self):
        """Initialize formatter"""
        self.doc = Document()
        self._setup_document_margins()

    def _setup_document_margins(self):
        """Set up document margins for better layout"""
        sections = self.doc.sections
        for section in sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

    def create_cornell_notes(self, notes_data: Dict, output_path: str):
        """
        Create Cornell Method notes document

        Args:
            notes_data: Structured notes data from AI
            output_path: Where to save the DOCX file
        """
        # Add title
        self._add_title(notes_data.get("title", "Study Notes"))

        # Add sections with Cornell layout
        sections = notes_data.get("sections", [])
        for section_idx, section in enumerate(sections):
            self._add_cornell_section(section, section_idx)

            # Add spacing between sections
            if section_idx < len(sections) - 1:
                self.doc.add_paragraph()

        # Add summary at the end
        if notes_data.get("summary"):
            self._add_summary(notes_data["summary"])

        # Save document
        output_file = Path(output_path)
        output_file.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(output_file))

    def _add_title(self, title: str):
        """
        Add formatted title to document

        Args:
            title: Document title
        """
        title_para = self.doc.add_paragraph()
        title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        title_run = title_para.add_run(title)
        title_run.font.size = FontSizes.TITLE
        title_run.font.name = FontNames.PRIMARY
        title_run.font.color.rgb = Colors.TITLE
        title_run.bold = True

        title_para.space_after = Pt(18)

        # Add decorative line
        self.doc.add_paragraph("_" * 80).alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    def _add_cornell_section(self, section: Dict, section_idx: int):
        """
        Add a Cornell Method section with two-column layout

        Args:
            section: Section data with cue_questions and main_notes
            section_idx: Section index for tracking
        """
        # Create table for Cornell layout (2 columns)
        table = self.doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'

        # Set column widths
        table.columns[0].width = CUE_COLUMN_WIDTH
        table.columns[1].width = NOTES_COLUMN_WIDTH

        # Style the table
        self._style_table(table)

        # Get cells
        cue_cell = table.rows[0].cells[0]
        notes_cell = table.rows[0].cells[1]

        # Fill cue questions (left column)
        cue_questions = section.get("cue_questions", [])
        self._add_cue_questions(cue_cell, cue_questions)

        # Fill main notes (right column)
        main_notes = section.get("main_notes", {})
        self._add_main_notes(notes_cell, main_notes)

    def _style_table(self, table):
        """
        Apply styling to Cornell layout table

        Args:
            table: python-docx table object
        """
        # Set border color
        tbl = table._element
        tblPr = tbl.tblPr if tbl.tblPr is not None else tbl.add_tblPr()

        # Create border elements
        tblBorders = OxmlElement('w:tblBorders')
        for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '8')  # Border width
            border.set(qn('w:color'), 'BFBFBF')  # Gray color
            tblBorders.append(border)

        tblPr.append(tblBorders)

    def _add_cue_questions(self, cell, questions: List[str]):
        """
        Add cue questions to left column

        Args:
            cell: Table cell for cue questions
            questions: List of cue questions
        """
        # Clear default paragraph
        cell.text = ""

        for idx, question in enumerate(questions):
            para = cell.add_paragraph()
            run = para.add_run(question)
            run.font.size = FontSizes.CUE_QUESTION
            run.font.name = FontNames.PRIMARY
            run.font.color.rgb = Colors.CUE_QUESTION
            run.italic = True

            # Add spacing between questions
            if idx < len(questions) - 1:
                para.space_after = Pt(12)

        # Set cell shading (light background)
        self._shade_cell(cell, Colors.SUMMARY_BG)

    def _add_main_notes(self, cell, notes: Dict):
        """
        Add main notes to right column with visual hierarchy

        Args:
            cell: Table cell for main notes
            notes: Dictionary with heading and content list
        """
        # Clear default paragraph
        cell.text = ""

        # Add section heading if present
        heading = notes.get("heading", "")
        if heading:
            heading_para = cell.add_paragraph()
            heading_run = heading_para.add_run(heading)
            heading_run.font.size = FontSizes.SECTION_HEADING
            heading_run.font.name = FontNames.PRIMARY
            heading_run.font.color.rgb = Colors.HEADING
            heading_run.bold = True
            heading_para.space_after = Pt(8)

        # Add content items
        content_items = notes.get("content", [])
        for item in content_items:
            self._add_content_item(cell, item)

    def _add_content_item(self, cell, item: Dict):
        """
        Add a single content item with appropriate formatting

        Args:
            cell: Table cell
            item: Content item dictionary with type, text, and level
        """
        item_type = item.get("type", "text")
        text = item.get("text", "")
        level = item.get("level", 0)

        if item_type == "heading":
            self._add_heading_item(cell, text, level)
        elif item_type == "bullet":
            self._add_bullet_item(cell, text, level)
        elif item_type == "emphasis":
            self._add_emphasis_item(cell, text, level)
        elif item_type == "table":
            self._add_table_item(cell, item)
        else:  # text
            self._add_text_item(cell, text, level)

    def _add_heading_item(self, cell, text: str, level: int):
        """Add a heading item"""
        para = cell.add_paragraph()
        run = para.add_run(text)
        run.font.size = Pt(12) if level == 1 else Pt(11)
        run.font.name = FontNames.PRIMARY
        run.font.color.rgb = Colors.HEADING
        run.bold = True

        # Indentation based on level
        para.paragraph_format.left_indent = Inches(0.25 * level)
        para.space_after = Pt(4)

    def _add_bullet_item(self, cell, text: str, level: int):
        """Add a bullet point item"""
        para = cell.add_paragraph(style='List Bullet')
        run = para.add_run(text)
        run.font.size = FontSizes.MAIN_TEXT
        run.font.name = FontNames.PRIMARY

        # Indentation based on level
        para.paragraph_format.left_indent = Inches(0.25 * level)
        para.space_after = Pt(3)

    def _add_emphasis_item(self, cell, text: str, level: int):
        """Add an emphasized item (key term/definition)"""
        para = cell.add_paragraph()
        run = para.add_run(text)
        run.font.size = FontSizes.MAIN_TEXT
        run.font.name = FontNames.PRIMARY
        run.bold = True

        # Add highlight
        run.font.highlight_color = 7  # Yellow highlight (WD_COLOR_INDEX.YELLOW)

        # Indentation based on level
        para.paragraph_format.left_indent = Inches(0.25 * level)
        para.space_after = Pt(4)

    def _add_text_item(self, cell, text: str, level: int):
        """Add a regular text item"""
        para = cell.add_paragraph()
        run = para.add_run(text)
        run.font.size = FontSizes.MAIN_TEXT
        run.font.name = FontNames.PRIMARY

        # Indentation based on level
        para.paragraph_format.left_indent = Inches(0.25 * level)
        para.space_after = Pt(4)

    def _add_table_item(self, cell, item: Dict):
        """Add a table within the notes"""
        headers = item.get("headers", [])
        rows = item.get("rows", [])

        if not headers or not rows:
            return

        # Add a small paragraph for spacing
        cell.add_paragraph().space_after = Pt(4)

        # Create nested table
        num_cols = len(headers)
        num_rows = len(rows) + 1  # +1 for header row

        nested_table = cell.add_table(rows=num_rows, cols=num_cols)
        nested_table.style = 'Light Grid Accent 1'

        # Add headers
        for col_idx, header in enumerate(headers):
            cell_obj = nested_table.rows[0].cells[col_idx]
            para = cell_obj.paragraphs[0]
            run = para.add_run(str(header))
            run.font.size = Pt(10)
            run.font.bold = True
            self._shade_cell(cell_obj, Colors.HEADING)

        # Add data rows
        for row_idx, row_data in enumerate(rows, start=1):
            for col_idx, cell_data in enumerate(row_data):
                if col_idx < num_cols:  # Safety check
                    cell_obj = nested_table.rows[row_idx].cells[col_idx]
                    para = cell_obj.paragraphs[0]
                    run = para.add_run(str(cell_data))
                    run.font.size = Pt(9)

        # Add spacing after table
        cell.add_paragraph().space_after = Pt(8)

    def _add_summary(self, summary_text: str):
        """
        Add summary section at the bottom

        Args:
            summary_text: Summary text
        """
        # Add spacing
        self.doc.add_paragraph()

        # Add "Summary" heading
        summary_heading = self.doc.add_paragraph()
        summary_heading_run = summary_heading.add_run("Summary")
        summary_heading_run.font.size = FontSizes.SECTION_HEADING
        summary_heading_run.font.name = FontNames.PRIMARY
        summary_heading_run.font.color.rgb = Colors.HEADING
        summary_heading_run.bold = True
        summary_heading.space_after = Pt(6)

        # Add summary text with background
        summary_para = self.doc.add_paragraph()
        summary_run = summary_para.add_run(summary_text)
        summary_run.font.size = FontSizes.SUMMARY
        summary_run.font.name = FontNames.PRIMARY

        # Add shading to summary paragraph
        pPr = summary_para._element.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'D6DCE4')  # Light blue background
        pPr.append(shd)

        # Add padding through spacing
        summary_para.paragraph_format.space_before = Pt(6)
        summary_para.paragraph_format.space_after = Pt(6)
        summary_para.paragraph_format.left_indent = Inches(0.25)
        summary_para.paragraph_format.right_indent = Inches(0.25)

    def _shade_cell(self, cell, color: RGBColor):
        """
        Add background shading to a table cell

        Args:
            cell: Table cell object
            color: RGBColor for background
        """
        cell_elem = cell._element
        cell_pr = cell_elem.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        # RGBColor converts to hex string automatically
        hex_color = str(color)  # Returns hex like 'FFFFFF'
        shd.set(qn('w:fill'), hex_color)
        cell_pr.append(shd)


def create_cornell_docx(notes_data: Dict, output_path: str):
    """
    Convenience function to create Cornell notes DOCX

    Args:
        notes_data: Structured notes data
        output_path: Output file path
    """
    formatter = CornellFormatter()
    formatter.create_cornell_notes(notes_data, output_path)
